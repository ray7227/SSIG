import re
from datetime import datetime, timedelta, time as dtime
from pathlib import Path
from zoneinfo import ZoneInfo

import io
import json
import math
import os
import re
import tempfile
import urllib.parse
import urllib.request
import zipfile
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, time as dtime
from difflib import get_close_matches
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
import streamlit as st

try:
    from astral import LocationInfo
    from astral.sun import sun
    ASTRAL_OK = True
except ImportError:
    ASTRAL_OK = False

# Free local OCR for the Photos tab (needs Tesseract installed on the machine).
try:
    import pytesseract
    from PIL import Image, ExifTags, ImageDraw, ImageFont
    from dateutil import parser as dateparser
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
    except Exception:
        pass
    OCR_LIBS = True
except Exception:
    OCR_LIBS = False

# Free offline dictionary for merging split words (optional).
try:
    from spellchecker import SpellChecker
    _SPELL = SpellChecker()
except Exception:
    _SPELL = None

# Optional GIS libs for the custom location picker.
try:
    import geopandas as gpd
    GEO_OK = True
except Exception:
    GEO_OK = False

try:
    import gpxpy
    GPX_OK = True
except Exception:
    GPX_OK = False

# =========================
# PAGE SETUP
# =========================
st.set_page_config(
    page_title="AiM Field Assistant",
    page_icon="🧭",
    layout="wide",
)
st.title("AiM Field Assistant")

SSIG_PDF = "https://open.alberta.ca/dataset/93d8a251-4a9a-428f-ad99-7484c6ebabe0/resource/f4024e81-b835-4a50-8fb1-5b31d9726b84/download/2013-sensitivespeciesinventoryguidelines-apr18.pdf"
SWEEP_PDF = "https://open.alberta.ca/dataset/d15221f2-f6d8-4671-8b49-d8fff6eab2b6/resource/6968392a-9e05-4bd8-bd76-ea107ba86c1c/download/aep-wildlife-sweep-protocols-sensitive-species-inventory-guidelines-2020.pdf"
FWMIS_CHECK = "https://www.alberta.ca/fwmis-loadform-check-tool"
FWMIS_GUIDE = "https://www.alberta.ca/system/files/custom_downloaded_images/ep-fwmis-data-submission-guide.pdf"

with st.sidebar:
    st.markdown("### Sources")
    st.markdown(f"[Sensitive Species Inventory Guidelines (2013)]({SSIG_PDF})")
    st.markdown(f"[Wildlife Sweep Protocols (2020)]({SWEEP_PDF})")

FILE = "SSIG_Breakdown.xlsx"
SWEEP_FILE = "Sweep_Breakdown.xlsx"
TZ = ZoneInfo("America/Edmonton")   # BC Peace users: MST year-round, adjust if needed

# Alberta location presets (lat, lon)
LOCATIONS = {
    "Calgary": (51.05, -114.07),
    "Edmonton": (53.55, -113.49),
    "Grande Prairie": (55.17, -118.80),
    "Medicine Hat": (50.04, -110.68),
    "Brooks": (50.58, -111.90),
    "Lethbridge": (49.69, -112.83),
    "Fort McMurray": (56.73, -111.38),
    "Fort St. John, BC": (56.25, -120.85),
    "Custom / map": None,
}

# =========================
# PHOTOS
# =========================
PHOTO_ROOT = "Images"

PHOTOS = {
    "Amphibians (Auditory Survey Guideline)": "Amph (Aud).png",
    "Amphibians (Non-Acoustic Survey Guideline)": "Amph (Non-ac).png",
    "Short-horned Lizard (ESHL)": "ESHL.png",
    "Snake Hibernacula Searches": "Snake.png",
    "Burrowing Owl (BUOW)": "BUOW.png",
    "Short-Eared Owl (SEOW) (BBS survey)": "SEOW.png",
    "Prairie Raptors (SR survey)": "SR.png",
    "Boreal & Foothills Raptors (BBS / SR survey)": "BBS & SR.png",
    "Grassland Birds (BBS survey)": "Grassland BBS.png",
    "Boreal & Foothills Breeding Songbirds & Woodpeckers (BBS survey)": "Boreal and foothills BBS.png",
    "Sharp-tailed Grouse (STGR)": "STGR.png",
    "Western Grebe (WEGR)": "WEGR.png",
    "Piping Plover (PIPL)": "PIPL.png",
    "Yellow Rail (BBS survey)": "Yellow Rail.png",
    "Common Nighthawk (CONI) (BBS survey)": "Nighthawk.png",
    "Bats": "Bats.png",
    "Swift Fox (SWFO)": "SWFO.png",
    "Ord's Kangaroo Rat (OKRA)": "OKRA.png",
    "Non-invasive Mammals Surveys-Winter Tracking & Searches for Mineral Licks": "Mineral licks.png",
    "Species at Risk Plant Surveys": "Plant.png",
}

def normalize(s):
    s = str(s).lower().strip()
    s = s.replace("\u2019", "'").replace("\u2018", "'")
    for dash in ("\u2013", "\u2014", "\u2011"):
        s = s.replace(dash, "-")
    s = re.sub(r"\s+", " ", s)
    return s

PHOTOS_NORM = {normalize(k): v for k, v in PHOTOS.items()}

def get_photo(survey):
    name = PHOTOS.get(survey) or PHOTOS_NORM.get(normalize(survey))
    if not name:
        return None
    path = Path(PHOTO_ROOT) / name
    return path if path.is_file() else None

# =========================
# LOAD SPREADSHEET
# =========================
@st.cache_data
def load_data(path):
    df = pd.read_excel(path, header=0)
    df = df.rename(columns={df.columns[0]: "Field"})
    df["Field"] = df["Field"].astype(str).str.strip()
    df = df.set_index("Field")
    df = df.dropna(axis=1, how="all").dropna(axis=0, how="all")
    return df

try:
    df = load_data(FILE)
except FileNotFoundError:
    st.error(f"Can't find {FILE}. Put it in the same folder as this app.")
    st.stop()

survey_types = list(df.columns)
fields = [f for f in df.index if f.lower() != "photo"]  # skip image row

def get_value(field, survey):
    val = df.loc[field, survey]
    if pd.isna(val):
        return ""
    return str(val).strip()

def field_value(field_query, survey):
    # match an index label by normalized name, exact first then contains
    for idx in df.index:
        if normalize(idx) == normalize(field_query):
            return get_value(idx, survey)
    for idx in df.index:
        if normalize(field_query) in normalize(idx):
            return get_value(idx, survey)
    return ""

try:
    sweep_df = load_data(SWEEP_FILE)
except FileNotFoundError:
    sweep_df = None

def render_browse(bdf, key, photo_fn=None):
    options = list(bdf.columns)
    flds = [f for f in bdf.index if f.lower() != "photo"]
    choice = st.selectbox("Select", options, key=key)
    st.header(choice)
    if photo_fn:
        p = photo_fn(choice)
        if p:
            left, right = st.columns([2, 3])
            with left:
                st.image(str(p), use_container_width=True)
            st.divider()
    for f in flds:
        val = bdf.loc[f, choice]
        if pd.isna(val) or not str(val).strip():
            continue
        st.markdown(f"**{f}**")
        st.write(str(val).strip())
        st.divider()

# =========================
# TIMING ENGINE
# =========================
def sun_times(lat, lon, d):
    loc = LocationInfo(latitude=lat, longitude=lon)
    s = sun(loc.observer, date=d, tzinfo=TZ)
    return s["sunrise"], s["sunset"]

def _clock(text):
    m = re.search(r"(\d{1,2})(?::(\d{2}))?\s*(am|pm)", text)
    if not m:
        return None
    h = int(m.group(1)) % 12
    if m.group(3) == "pm":
        h += 12
    return dtime(h, int(m.group(2) or 0))

def _offset(text, anchor):
    # returns a timedelta if <anchor> is present, else None. Zero if no number given.
    if anchor not in text:
        return None
    m = re.search(r"(\d+(?:\.\d+)?)\s*(hours?|hrs?|minutes?|mins?)\s+" + anchor, text)
    if not m:
        return timedelta(0)
    n = float(m.group(1))
    unit = m.group(2)
    return timedelta(hours=n) if unit.startswith(("hour", "hr")) else timedelta(minutes=n)

def compute_window(rule, sunrise, sunset, d):
    """Best-effort start/end from a Time-of-Day rule. Returns (start, end)."""
    t = normalize(rule)
    start = end = None

    off = _offset(t, "before sunrise")
    if off is not None:
        start = sunrise - off
    elif (off := _offset(t, "after sunrise")) is not None:
        start = sunrise + off
    elif (off := _offset(t, "after sunset")) is not None:
        start = sunset + off

    off = _offset(t, "before sunset")
    if off is not None:
        end = sunset - off

    if ("no later than" in t) or ("until" in t):
        tail = t.split("no later than")[-1] if "no later than" in t else t.split("until")[-1]
        c = _clock(tail)
        if c:
            end = datetime.combine(d, c).replace(tzinfo=TZ)
            if start and end <= start:
                end += timedelta(days=1)

    if start is None and end is None and "daylight" in t:
        start, end = sunrise, sunset

    return start, end

def fmt(dt):
    return dt.strftime("%H:%M") if dt else "-"

# =========================
# SAFETY FORMS
# =========================
def safety_forms(vehicle, first_day, water_ice, prime, drive_hr, staff_type):
    flha_when = "Daily (Prime Contractor)" if prime else "Daily"
    submit = [("FLHA / Tailgate", flha_when, "SafetyAdmin")]
    if first_day:
        submit.append(("Project Orientation / Kickoff Meeting", "First day", "SafetyAdmin"))
        submit.append(("Project ERP + First Aid Assessment", "First day", "SafetyAdmin"))
    if vehicle == "AiM-owned":
        submit.append(("Vehicle inspection", "Weekly", "Fleetio"))
    elif vehicle == "Personal":
        submit.append(("Vehicle inspection", "Monthly", "SafetyAdmin"))
    elif vehicle == "Rental":
        submit.append(("Vehicle inspection", "Daily", "Fleetio"))
    if water_ice == "Water":
        submit.append(("Working on Water", "Before water work", "SafetyAdmin"))
    elif water_ice == "Ice":
        submit.append(("Working on Ice + Ice Rod", "Before ice work", "SafetyAdmin"))
    if drive_hr > 3:
        submit.append(("Journey Mgmt Plan", "Drive over 3 hr", "SafetyAdmin"))
    hazard_when = "Monthly (field staff)" if staff_type == "Field" else "Yearly (remote)"
    submit.append(("Hazard ID", hazard_when, "SafetyAdmin"))
    as_needed = "Near Miss, Incident"
    return submit, as_needed

# =========================
# WILDLIFE SWEEP REFERENCE (AEP 2020 protocol, paraphrased)
# =========================
SWEEP_INFO = [
    ("What it is", "A walkthrough of the proposed disturbance site plus a surrounding buffer to find important wildlife features that must be avoided. Required before land or vegetation disturbance on applicable dispositions."),
    ("Features to find", "Occupied raptor nests, heron rookeries, occupied dens, hibernacula, and natural mineral licks. Also osprey and bald eagle nests outside the Grassland and Parkland regions at any time of year."),
    ("Coverage", "The disturbance footprint plus a 100 m buffer."),
    ("Timing", "Within 10 days before activity starts, as close to day one as possible. Earlier is allowed with justification (e.g. bear dens before deep snow)."),
    ("Conditions", "Daylight, and weather that does not obscure detection."),
    ("Personnel", "Someone with the education, knowledge, or experience to locate and identify wildlife features. No wildlife research permit needed."),
    ("If a feature is found", "Photograph and georeference it without flushing the occupant. Avoid it by re-siting or re-timing, or submit a non-routine application with justification and mitigation. If the disposition is already issued and the feature sits within a setback, contact the Regulator before entry."),
    ("If nothing is found", "Keep the sweep record for the life of the disposition. Activity may proceed subject to other requirements."),
    ("Records to keep", "Timing and its justification, personnel, environmental conditions, a GPS track of the walkthrough, and photos plus coordinates of any features. Held by the disposition holder for the disposition duration."),
    ("Occupied nest", "In current use: a bird present, territorial display, fresh feces, and/or feathers. Raptor nests stay active through the next year, with active status dropped June 1 of the second year of inactivity."),
    ("Occupied den", "In current use: an animal present, territorial display, fresh feces, digging or excavation, and/or tracks. Occupancy varies by season (breeding vs hibernation den)."),
    ("Liability", "Rests with the Disposition Holder, Licensee, or Permit Holder. An inadequate sweep that leads to disturbance can bring enforcement under the Wildlife Act or Public Lands Act."),
    ("Exemptions", "Dispositions in an EPEA-approved area with an authorized WMMP whose sweep protocol meets or exceeds this one."),
    ("Does not replace", "A full wildlife survey where one is required, disposition approval conditions, or federal MBCA/SARA requirements (contact ECCC for those)."),
]

# =========================
# CLUBROOT SAMPLING REFERENCE (AiM SOP AB_C_E002, condensed)
# =========================
CLUBROOT_INFO = [
    ("What it is", "Soil sampling to detect clubroot (Plasmodiophora brassicae), a soil-borne disease of canola and other crucifers. Spores travel on soil stuck to equipment, tires, and boots, so results set the cleaning requirements during construction."),
    ("Where to sample", "The construction footprint and the primary agricultural access point of each cultivated field. On pipeline projects, land is split into ownership-based tracts (usually quarter sections)."),
    ("Footprint sampling", "Walk the footprint (incl. ROW and temp workspace) in a zigzag. Take ~100 g topsoil at even spacing from both edges and the center, prioritizing low or wet spots. 5 samples per tract combined into one composite. A quarter section is ~800 m, so about every 150 m."),
    ("Ag access sampling", "5 samples in a W pattern at the primary agricultural access of each tract. Run alongside footprint sampling, or do footprint first then access points if they're far off (800+ m)."),
    ("Sample handling", "Label and seal bags, store in coolers, deliver to the lab as soon as possible (confirm the lab with the PM)."),
    ("Equipment cleaning", "Required at every new tract or change in land ownership. Assume every field may be contaminated. Brush and pick off soil first, then spray boots, shovel, and gear with 2.0% bleach solution, brush, and mist again. Log the Clubroot Cleaning Form in SafetyAdmin with geotagged, timestamped photos. Boot covers are an option but tear easily (carry two sets)."),
    ("Cleaning kit", "2.0% bleach solution (or equivalent), spray bottle(s), brush(es), pick(s), boot covers (optional)."),
    ("Pre-planning", "Build a clubroot tracker in Excel (legal location, tract ID, access). Use Google Earth plus a KML of the footprint and tracts (from GIS). Access only off public roads; no crossing private land without landowner permission. Note In-and-Out vs Walk-through tracts and leap-frog vehicles. Load KMZ/KML to GPS or Avenza."),
    ("Weather", "Do not sample in adverse weather."),
    ("Forms", "FLHA Tailgate, Vehicle Inspection, Journey Mgmt Plan, POKA (kickoff), ERP + First Aid, Prime Contractor Tailgate if applicable, plus the project Clubroot Cleaning Form."),
    ("After / back home", "Complete OkAlone/SPOT/InReach check-outs. Upload cleaning forms, GPS tracks, waypoints, and photos to the SharePoint project folder, named FirstInitialLastName_Clubroot_Date. Send forms plus zipped data to the PM. When lab results arrive, update the tracker and assign a risk level per spores/gram, which sets future cleaning requirements."),
    ("Key hazards", "Uneven terrain, weather, landowners and dogs, barbed wire, bleach, fatigue, working alone, wildlife, insect bites, farm and construction equipment."),
]

# =========================
# EXPENSES REFERENCE (Vantage Point)
# =========================
EXPENSES_INFO = [
    ("Daily meals", "Meals", "$20 x 3 meals = $60/day"),
    ("Truck mileage (actual)", "Mileage - Trucks", "Enter the day's km"),
    ("Fuel", "Mileage - flat rate", "Attach the fuel receipt to get reimbursed"),
    ("Data access fee", "Data access fee", "$25"),
]

# =========================
# OCR HELPERS (free, local Tesseract)
# =========================
_MON = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"
DATE_RE = re.compile(
    r"(20\d{2}[-/.\s]\d{1,2}[-/.\s]\d{1,2})"
    r"|(\d{1,2}[-/.\s]\d{1,2}[-/.\s]20\d{2})"
    r"|(" + _MON + r"\s+\d{1,2},?\s+20\d{2})"
    r"|(\d{1,2}\s+" + _MON + r"\s+20\d{2})",
    re.IGNORECASE,
)

def _tesseract_ok():
    if not OCR_LIBS:
        return False
    try:
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False

def _open_image(data):
    return Image.open(io.BytesIO(data))

def _to_jpeg_bytes(data):
    img = _open_image(data)
    if img.mode != "RGB":
        img = img.convert("RGB")
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=90)
    return buf.getvalue()

@st.cache_data(show_spinner=False)
def _jpeg_cached(data):
    try:
        return _to_jpeg_bytes(data)
    except Exception:
        return data

def _exif_date(data):
    try:
        img = _open_image(data)
        exif = img.getexif()
        raw = exif.get(306)  # DateTime
        if not raw:
            ifd = exif.get_ifd(0x8769)  # Exif IFD
            raw = ifd.get(36867)  # DateTimeOriginal
        if raw:
            return datetime.strptime(str(raw)[:10], "%Y:%m:%d").strftime("%Y-%m-%d")
    except Exception:
        pass
    return ""

def _dms(vals, ref):
    if not vals:
        return None
    d, m, s = [float(x) for x in vals]
    v = d + m / 60 + s / 3600
    return -v if ref in ("S", "W") else v

def _exif_gps(data):
    try:
        exif = _open_image(data).getexif()
        gps = exif.get_ifd(0x8825)
        if not gps:
            return None
        lat = _dms(gps.get(2), gps.get(1))
        lon = _dms(gps.get(4), gps.get(3))
        if lat is None or lon is None:
            return None
        return lat, lon
    except Exception:
        return None

def _dist_m(a, b):
    R = 6371000
    lat1, lon1 = map(math.radians, a)
    lat2, lon2 = map(math.radians, b)
    dlat, dlon = lat2 - lat1, lon2 - lon1
    x = math.sin(dlat / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin(dlon / 2) ** 2
    return 2 * R * math.asin(math.sqrt(x))

TIME_RE = re.compile(r"\d{1,2}:\d{1,2}(?::\d{1,2})?(?:\s*[-+]\d{1,2}:?\d{2})?")

def _parse_date(text):
    m = DATE_RE.search(text)
    if not m:
        return ""
    try:
        return dateparser.parse(m.group(0), dayfirst=False).strftime("%Y-%m-%d")
    except Exception:
        return ""

def _label_from_text(text):
    t = re.sub(r"\s+", " ", text)
    t = DATE_RE.sub("", t)
    t = TIME_RE.sub("", t)
    t = re.sub(r"[^A-Za-z0-9 _-]", "", t)
    t = re.sub(r"\s+", "_", t).strip("_").lower()
    return t[:40] or "photo"

def _prep(im, thr=205):
    # Sharpen white banner text: grayscale, threshold to black-on-white.
    # Only upscale small crops; big phone-photo crops are already sharp enough.
    g = im.convert("L")
    if g.width < 1200:
        g = g.resize((g.width * 2, g.height * 2))
    return g.point(lambda p: 0 if p > thr else 255)

@st.cache_data(show_spinner=False)
def ocr_photo(data: bytes, mode: str = "banner", band: float = 0.08):
    img = _open_image(data)
    if img.mode != "RGB":
        img = img.convert("RGB")
    w, h = img.size
    if mode == "banner":
        top = int(h * (1 - band))
        label_img = img.crop((0, top, int(w * 0.66), h))       # left + middle
        label = _label_from_text(
            pytesseract.image_to_string(_prep(label_img, 205), config="--psm 6")
        )
        date = _exif_date(data)
        if not date:
            date_img = img.crop((int(w * 0.66), top, w, h))    # right corner
            date_txt = pytesseract.image_to_string(
                _prep(date_img, 185),
                config="--psm 6 -c tessedit_char_whitelist=0123456789-:",
            )
            date = _parse_date(date_txt)
    else:
        text = pytesseract.image_to_string(img)
        label = _label_from_text(text)
        date = _exif_date(data) or _parse_date(text)
    return label, date

def _safe(name):
    name = re.sub(r"[^A-Za-z0-9_-]+", "_", str(name).strip()).strip("_").lower()
    return name or "photo"

def _is_word(w):
    if _SPELL is None or len(w) < 3:
        return False
    return len(_SPELL.known([w])) == 1

def _merge_tokens(tokens):
    # Rejoin words OCR split apart, e.g. ['p', 'lacement'] -> ['placement'].
    out = []
    i = 0
    while i < len(tokens):
        t = tokens[i]
        if i + 1 < len(tokens):
            joined = t + tokens[i + 1]
            if _is_word(joined) and (not _is_word(t) or not _is_word(tokens[i + 1])):
                out.append(joined)
                i += 2
                continue
        out.append(t)
        i += 1
    return out

def clean_labels(labels):
    # 1) Rejoin split words. 2) Fix garbled tokens against the batch vocabulary.
    merged = ["_".join(_merge_tokens([t for t in l.split("_") if t])) for l in labels]
    toks = []
    for l in merged:
        toks += [t for t in l.split("_") if t]
    counts = Counter(toks)
    n = len(labels)
    thresh = max(2, round(n * 0.2))
    common = [t for t, c in counts.items() if c >= thresh and len(t) >= 2 and not t.isdigit()]
    cset = set(common)
    out = []
    for l in merged:
        parts = []
        for t in l.split("_"):
            if not t:
                continue
            if t in cset or len(t) < 2 or t.isdigit():
                parts.append(t)
                continue
            m = get_close_matches(t, common, n=1, cutoff=0.8)
            parts.append(m[0] if m else t)
        out.append("_".join(parts) or "photo")
    return out

# =========================
# WEATHER (free, Open-Meteo, no key)
# =========================
WMO = {
    0: "Clear", 1: "Mainly clear", 2: "Partly cloudy", 3: "Overcast",
    45: "Fog", 48: "Fog",
    51: "Drizzle", 53: "Drizzle", 55: "Drizzle", 56: "Freezing drizzle", 57: "Freezing drizzle",
    61: "Rain", 63: "Rain", 65: "Heavy rain", 66: "Freezing rain", 67: "Freezing rain",
    71: "Snow", 73: "Snow", 75: "Heavy snow", 77: "Snow grains",
    80: "Rain showers", 81: "Rain showers", 82: "Heavy showers",
    85: "Snow showers", 86: "Snow showers",
    95: "Thunderstorm", 96: "Thunderstorm", 99: "Thunderstorm",
}
RAIN_CODES = set(range(51, 68)) | set(range(80, 83)) | {95, 96, 99}

@st.cache_data(show_spinner=False, ttl=3600)
def fetch_weather(lat, lon, start, end):
    params = {
        "latitude": lat,
        "longitude": lon,
        "daily": "weather_code,temperature_2m_max,temperature_2m_min,precipitation_sum,wind_speed_10m_max",
        "timezone": "America/Edmonton",
        "start_date": start,
        "end_date": end,
    }
    url = "https://api.open-meteo.com/v1/forecast?" + urllib.parse.urlencode(params)
    with urllib.request.urlopen(url, timeout=20) as resp:
        return json.loads(resp.read().decode())

def parse_restrictions(text):
    t = (text or "").lower()
    wind = None
    m = re.search(r"(\d{1,3})\s*km", t)
    if m:
        wind = int(m.group(1))
    temp_min = None
    m = re.search(r"(?:≥|>=|at least|above)\s*(\d{1,2})\s*°?\s*c", t)
    if not m:
        m = re.search(r"(\d{1,2})\s*°?\s*c", t)
    if m:
        temp_min = int(m.group(1))
    no_rain = ("no rain" in t) or ("avoid" in t and "rain" in t)
    return wind, temp_min, no_rain

# =========================
# PLAN BUILDER + WORD EXPORT
# =========================
try:
    from docx import Document
    DOCX_OK = True
except Exception:
    DOCX_OK = False

def _sub(text): return {"type": "sub", "text": text}
def _bul(items): return {"type": "bul", "items": [i for i in items if i]}
def _par(text): return {"type": "par", "text": text}
def _tab(cols, rows): return {"type": "tab", "cols": cols, "rows": rows}

def build_plan(work, survey, loc_name, lat, lon, start, end,
               vehicle, staff_type, water_ice, drive_hr, first_day, prime):
    before, during, after = [], [], []
    sr = ss = None
    is_survey = (work == "Protocol survey")
    restr = field_value("Weather Restrictions", survey) if is_survey else ""
    method = field_value("Method", survey) if is_survey else ""
    crew = field_value("Required Survey Personnel", survey) if is_survey else ""
    window = field_value("Survey Window", survey) if is_survey else ""
    tod = field_value("Time of Day", survey) if is_survey else ""

    submit, _ = safety_forms(vehicle, first_day, water_ice, prime, drive_hr, staff_type)
    before_forms, during_forms = [], []
    for name, when, where in submit:
        if when == "First day" or name.startswith("Journey") or name.startswith("Hazard"):
            before_forms.append((name, when, where))
        else:
            during_forms.append((name, when, where))

    # ===== BEFORE =====
    before.append(_sub("Weather"))
    if ASTRAL_OK:
        try:
            sr, ss = sun_times(lat, lon, start)
            before.append(_par(f"{start:%b %d}: sunrise {fmt(sr)}, sunset {fmt(ss)} ({loc_name})."))
        except Exception:
            sr = ss = None
    try:
        wx = fetch_weather(lat, lon, start.isoformat(), end.isoformat())
        wind_limit, temp_min, no_rain = parse_restrictions(restr)
        d = wx["daily"]
        rows = []
        for i, day in enumerate(d["time"]):
            code = d["weather_code"][i]; hi = d["temperature_2m_max"][i]; lo = d["temperature_2m_min"][i]
            wind = d["wind_speed_10m_max"][i]; rain = d["precipitation_sum"][i]
            flags = []
            if wind_limit and wind and wind > wind_limit: flags.append("wind")
            if temp_min and hi is not None and hi < temp_min: flags.append("cold")
            if no_rain and (code in RAIN_CODES or (rain and rain >= 1)): flags.append("rain")
            rows.append([day, WMO.get(code, "?"),
                         round(hi) if hi is not None else "-",
                         round(lo) if lo is not None else "-",
                         round(wind) if wind is not None else "-",
                         round(rain, 1) if rain is not None else "-",
                         ", ".join(flags) if flags else "OK"])
        before.append(_tab(["Date", "Cond", "Hi", "Lo", "Wind", "Rain", "Flag"], rows))
        lim = []
        if wind_limit: lim.append(f"wind <= {wind_limit} km/h")
        if temp_min: lim.append(f"temp >= {temp_min} C")
        if no_rain: lim.append("avoid rain")
        if lim: before.append(_par("Survey limits: " + ", ".join(lim) + "."))
    except Exception:
        before.append(_par("Forecast unavailable (needs internet; reaches about 16 days out)."))

    if is_survey and tod:
        before.append(_sub("Survey window"))
        before.append(_bul([f"Window: {window}" if window else "", f"Time of day: {tod}"]))

    before.append(_sub("Before you go"))
    prep = [f"{n} ({wn}, {wr})" for n, wn, wr in before_forms]
    prep += ["Confirm access, permits, and landowner contact.",
             "Set up check-in/out (OkAlone / SPOT / InReach)."]
    if work == "Clubroot sampling":
        prep.append("Pack cleaning kit: 2.0% bleach, spray bottle, brushes, picks, boot covers.")
        prep.append("Build the tract tracker and load KMZ/KML to GPS or Avenza.")
    if is_survey:
        prep.append("Review species ID and method in the Reference tab.")
    before.append(_bul(prep))

    # ===== DURING =====
    during.append(_sub("The work"))
    if is_survey:
        wl = [f"Survey: {survey}", f"Method: {method or '-'}", f"Crew: {crew or '-'}"]
        if ASTRAL_OK and tod and sr and ss:
            s0, e0 = compute_window(tod, sr, ss, start)
            if s0 and e0:
                wl.append(f"Timing (start day): {fmt(s0)}-{fmt(e0)}")
        during.append(_bul(wl))
    elif work == "Wildlife sweep":
        during.append(_bul([
            "Walk the footprint plus a 100 m buffer, in daylight.",
            "Watch for occupied nests, dens, hibernacula, and mineral licks.",
            "Photograph and GPS any feature without flushing it.",
        ]))
    else:
        during.append(_bul([
            "Clean at every new tract or ownership change (2.0% bleach).",
            "Footprint: zigzag, ~100 g, 5 per tract, edges + center, low/wet spots.",
            "Ag access: 5 in a W pattern. Label, seal, cooler.",
        ]))
    during.append(_sub("Daily forms"))
    during.append(_bul([f"{n} ({wn}, {wr})" for n, wn, wr in during_forms]))
    if is_survey and restr:
        during.append(_sub("Watch-outs"))
        during.append(_par(restr))

    # ===== AFTER =====
    after.append(_sub("Records"))
    recs = ["Complete OkAlone / SPOT / InReach check-out.",
            "Upload GPS tracks, waypoints, photos, and forms to the SharePoint project folder.",
            "Name files FirstInitialLastName_Activity_Date."]
    if work == "Clubroot sampling":
        recs.append("Deliver samples to the lab, update the tracker, assign risk from lab results.")
    after.append(_bul(recs))
    if is_survey or work == "Wildlife sweep":
        after.append(_sub("FWMIS"))
        after.append(_bul(["Compile every species observed into the FWMIS loadform.",
                           f"Check tool: {FWMIS_CHECK}", f"Guide: {FWMIS_GUIDE}"]))
    after.append(_sub("Photos"))
    after.append(_par("Use the Photos tab to OCR-label and sort field photos by date."))
    after.append(_sub("Expenses"))
    after.append(_tab(["Expense", "Vantage Point", "Amount / notes"], [list(r) for r in EXPENSES_INFO]))

    return [("Before", before), ("During", during), ("After", after)]

def render_blocks(blocks):
    for b in blocks:
        if b["type"] == "sub":
            st.markdown(f"**{b['text']}**")
        elif b["type"] == "bul":
            for it in b["items"]:
                st.markdown(f"- {it}")
        elif b["type"] == "par":
            st.write(b["text"])
        elif b["type"] == "tab":
            st.table(pd.DataFrame(b["rows"], columns=b["cols"]))

def build_docx(title, header_lines, sections):
    doc = Document()
    doc.add_heading(title, level=0)
    for l in header_lines:
        if l:
            doc.add_paragraph(l)
    for name, blocks in sections:
        doc.add_heading(name, level=1)
        for b in blocks:
            if b["type"] == "sub":
                doc.add_heading(b["text"], level=2)
            elif b["type"] == "par":
                doc.add_paragraph(b["text"])
            elif b["type"] == "bul":
                for it in b["items"]:
                    doc.add_paragraph(it, style="List Bullet")
            elif b["type"] == "tab":
                t = doc.add_table(rows=1, cols=len(b["cols"]))
                try:
                    t.style = "Light Grid Accent 1"
                except Exception:
                    pass
                for j, c in enumerate(b["cols"]):
                    t.rows[0].cells[j].text = str(c)
                for r in b["rows"]:
                    cells = t.add_row().cells
                    for j, v in enumerate(r):
                        cells[j].text = str(v)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# =========================
# LOCATION PICKER (place search, shapefile, map)
# =========================
def gpx_points(data_bytes):
    g = gpxpy.parse(data_bytes.decode("utf-8", "ignore"))
    pts = [(p.latitude, p.longitude) for t in g.tracks for s in t.segments for p in s.points]
    for r in g.routes:
        pts += [(p.latitude, p.longitude) for p in r.points]
    wpts = [(w.latitude, w.longitude, w.name or "") for w in g.waypoints]
    return pts, wpts

def _load_font(size):
    for p in ["/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", "DejaVuSans-Bold.ttf", "Arial.ttf"]:
        try:
            return ImageFont.truetype(p, size)
        except Exception:
            continue
    return ImageFont.load_default()

def _outline(draw, xy, text, font):
    x, y = xy
    for dx in (-2, -1, 0, 1, 2):
        for dy in (-2, -1, 0, 1, 2):
            if dx or dy:
                draw.text((x + dx, y + dy), text, font=font, fill="black")
    draw.text((x, y), text, font=font, fill="white")

def caption_photo(data, caption, coords="", when=""):
    img = _open_image(data).convert("RGB")
    W, H = img.size
    draw = ImageDraw.Draw(img)
    fs = max(20, W // 40)
    font = _load_font(fs)
    pad = fs
    if caption:
        _outline(draw, (pad, H - pad - fs), caption, font)
    right = [l for l in [coords, when] if l]
    y = H - pad - fs * len(right) - (len(right) - 1) * 4 if right else 0
    for line in right:
        w = draw.textlength(line, font=font)
        _outline(draw, (W - pad - w, y), line, font)
        y += fs + 4
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=90)
    return buf.getvalue()

def _photon_raw(q):
    params = {"q": q, "limit": 10, "lang": "en", "lat": 54.5, "lon": -114.5}
    url = "https://photon.komoot.io/api/?" + urllib.parse.urlencode(params)
    req = urllib.request.Request(url, headers={"User-Agent": "aim-field-assistant"})
    with urllib.request.urlopen(req, timeout=10) as r:
        return json.loads(r.read().decode()).get("features") or []

def _geo_openmeteo(q):
    url = "https://geocoding-api.open-meteo.com/v1/search?" + urllib.parse.urlencode(
        {"name": q, "count": 10, "language": "en", "format": "json"}
    )
    with urllib.request.urlopen(url, timeout=10) as r:
        data = json.loads(r.read().decode())
    res = data.get("results") or []
    if not res:
        raise ValueError("no result")
    ab = [x for x in res if x.get("country_code") == "CA" and x.get("admin1") == "Alberta"]
    ca = [x for x in res if x.get("country_code") == "CA"]
    pick = (ab or ca or res)[0]
    return pick["latitude"], pick["longitude"]

@st.cache_data(show_spinner=False, ttl=86400)
def geocode_place(q):
    ql = q.strip()
    has_region = bool(re.search(r"alberta|canada|\b(ab|bc|sk|mb|on)\b", ql, re.I))
    tries = ([ql + ", Alberta, Canada", ql + ", Canada"] if not has_region else []) + [ql]
    for t in tries:
        try:
            feats = _photon_raw(t)
        except Exception:
            feats = []
        if not feats:
            continue
        feats.sort(
            key=lambda f: (f.get("properties", {}).get("state") == "Alberta") * 2
            + (f.get("properties", {}).get("countrycode") == "CA"),
            reverse=True,
        )
        top = feats[0]
        if not has_region and top.get("properties", {}).get("countrycode") != "CA":
            continue
        lon, lat = top["geometry"]["coordinates"][:2]
        return lat, lon
    return _geo_openmeteo(ql)

def shapefile_centroid(uploaded):
    data = uploaded.getvalue()
    name = uploaded.name.lower()
    with tempfile.TemporaryDirectory() as tmp:
        if name.endswith(".zip"):
            p = os.path.join(tmp, "in.zip")
            with open(p, "wb") as fh:
                fh.write(data)
            gdf = gpd.read_file("zip://" + p)
        else:
            ext = os.path.splitext(name)[1] or ".geojson"
            p = os.path.join(tmp, "in" + ext)
            with open(p, "wb") as fh:
                fh.write(data)
            gdf = gpd.read_file(p)
    if gdf.crs is None:
        gdf = gdf.set_crs(4326, allow_override=True)
    gdf = gdf.to_crs(4326)
    geom = gdf.geometry.union_all() if hasattr(gdf.geometry, "union_all") else gdf.geometry.unary_union
    c = geom.centroid
    return float(c.y), float(c.x)

def location_picker(key="loc"):
    ss = st.session_state
    pk_lat, pk_lon, tk = f"{key}_platlat", f"{key}_platlon", f"{key}_track"
    ss.setdefault(pk_lat, 51.05)
    ss.setdefault(pk_lon, -114.07)
    ss.setdefault(tk, [])

    up_types = (["gpx"] if GPX_OK else []) + (["zip", "geojson", "json", "kml"] if GEO_OK else [])
    if up_types:
        shp = st.file_uploader(
            "Drop a GPX, shapefile (.zip), GeoJSON, or KML to set the location",
            type=up_types,
            key=f"{key}_shp",
        )
        if shp is not None:
            try:
                if shp.name.lower().endswith(".gpx"):
                    pts, wpts = gpx_points(shp.getvalue())
                    ss[tk] = pts
                    ss["gpx_waypoints"] = wpts
                    allp = pts + [(a, b) for a, b, _ in wpts]
                    if allp:
                        ss[pk_lat] = sum(p[0] for p in allp) / len(allp)
                        ss[pk_lon] = sum(p[1] for p in allp) / len(allp)
                        st.caption(f"{len(pts)} track points, {len(wpts)} waypoints loaded.")
                else:
                    la, lo = shapefile_centroid(shp)
                    ss[tk] = []
                    ss["gpx_waypoints"] = []
                    ss[pk_lat], ss[pk_lon] = la, lo
                    st.caption(f"Centroid from file: {la:.5f}, {lo:.5f}")
            except Exception as e:
                st.warning(f"Couldn't read that file: {e}")
    else:
        st.caption("Install geopandas or gpxpy to drop GIS files here.")

    q = st.text_input("Or search a place", key=f"{key}_q", placeholder="e.g. Grande Prairie, AB")
    if st.button("Find place", key=f"{key}_find") and q.strip():
        try:
            la, lo = geocode_place(q.strip())
            ss[pk_lat], ss[pk_lon] = la, lo
            ss[tk] = []
            ss["gpx_waypoints"] = []
        except Exception as e:
            st.warning(f"Not found: {e}")

    lat = st.number_input("Latitude", value=float(ss[pk_lat]), format="%.5f")
    lon = st.number_input("Longitude", value=float(ss[pk_lon]), format="%.5f")
    ss[pk_lat], ss[pk_lon] = lat, lon

    map_df = pd.DataFrame(
        [{"lat": a, "lon": b} for a, b in ss[tk]] + [{"lat": lat, "lon": lon}]
    )
    try:
        st.map(map_df, zoom=10)
    except Exception:
        st.map(map_df)
    return lat, lon

# =========================
# UI
# =========================
tab_plan, tab_ref, tab_photos = st.tabs(["Field Plan", "Reference", "Photos"])

# ---- FIELD PLAN ----
with tab_plan:
    c1, c2 = st.columns(2)
    with c1:
        work = st.selectbox("Work type", ["Protocol survey", "Wildlife sweep", "Clubroot sampling"])
    survey = None
    if work == "Protocol survey":
        with c2:
            survey = st.selectbox("Survey", survey_types)

    loc_name = st.selectbox("Location", list(LOCATIONS.keys()))
    if LOCATIONS[loc_name] is None:
        lat, lon = location_picker("planloc")
    else:
        lat, lon = LOCATIONS[loc_name]

    d1, d2 = st.columns(2)
    with d1:
        start = st.date_input("Start")
    with d2:
        end = st.date_input("End")

    with st.expander("Safety details"):
        s1, s2 = st.columns(2)
        with s1:
            vehicle = st.selectbox("Vehicle", ["AiM-owned", "Personal", "Rental", "None"])
            staff_type = st.selectbox("Staff type", ["Field", "Remote"])
            water_ice = st.selectbox("Water / ice work", ["None", "Water", "Ice"])
        with s2:
            drive_hr = st.number_input("Drive one-way (hr)", min_value=0.0, value=0.0, step=0.5)
            first_day = st.checkbox("First day of project")
            prime = st.checkbox("AiM is Prime Contractor")

    if end < start:
        st.warning("End date is before start date.")
    else:
        sections = build_plan(work, survey, loc_name, lat, lon, start, end,
                              vehicle, staff_type, water_ice, drive_hr, first_day, prime)
        title = f"Field Plan - {survey or work}"
        header = [f"Work: {work}", f"Survey: {survey}" if survey else "",
                  f"Location: {loc_name}", f"Dates: {start:%b %d} to {end:%b %d, %Y}"]

        if DOCX_OK:
            st.download_button(
                "Download Word plan",
                build_docx(title, header, sections),
                file_name=f"field_plan_{start.isoformat()}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        else:
            st.caption("Add python-docx to requirements.txt to enable the Word download.")

        tb, td, ta = st.tabs(["Before", "During", "After"])
        for tab, (name, blocks) in zip([tb, td, ta], sections):
            with tab:
                render_blocks(blocks)

# ---- REFERENCE ----
with tab_ref:
    ref = st.radio("Reference", ["Survey", "Sweep", "Clubroot", "Search"], horizontal=True)
    if ref == "Survey":
        render_browse(df, "ref_survey", get_photo)
    elif ref == "Sweep":
        if sweep_df is not None:
            render_browse(sweep_df, "ref_sweep")
        else:
            st.table(pd.DataFrame(SWEEP_INFO, columns=["", "Detail"]))
            st.caption("Source: Wildlife Sweep Protocols 2020 (sidebar).")
    elif ref == "Clubroot":
        st.table(pd.DataFrame(CLUBROOT_INFO, columns=["", "Detail"]))
        st.caption("Source: AiM Clubroot Sampling SOP (AB_C_E002).")
    else:
        query = st.text_input("Search all guidelines", placeholder="wind, sunset, egg searches, transect...")
        if query.strip():
            q = query.strip().lower()
            hits = []
            for field in fields:
                for s in survey_types:
                    value = get_value(field, s)
                    if value and q in value.lower():
                        hits.append((s, field, value))
            if hits:
                st.caption(f"{len(hits)} match(es)")
                for s, field, value in hits:
                    st.markdown(f"**{field}  ·  {s}**")
                    st.write(value)
                    st.divider()
            else:
                st.info("No matches.")

# ---- PHOTOS ----
with tab_photos:
    st.subheader("Photos")

    if OCR_LIBS:
        with st.expander("Caption a photo (type text onto the bottom-left)"):
            cf = st.file_uploader("Photo", type=["jpg", "jpeg", "png", "heic", "heif"], key="cap_up")
            cap = st.text_input("Bottom-left caption", key="cap_text")
            cco = st.text_input("Bottom-right line 1 (e.g. 50.934N 113.963W)", key="cap_coord")
            cwh = st.text_input("Bottom-right line 2 (e.g. date)", key="cap_when")
            if cf is not None and st.button("Add caption", key="cap_btn"):
                out = caption_photo(cf.getvalue(), cap.strip(), cco.strip(), cwh.strip())
                st.image(out)
                st.download_button(
                    "Download captioned photo", out,
                    file_name="captioned.jpg", mime="image/jpeg", key="cap_dl",
                )
    st.divider()

    st.caption("Reads the label in each photo, pre-fills a name and date, you fix any misses, then download a zip sorted by date. Free, runs on your machine.")

    if not OCR_LIBS:
        st.info("Add pillow, pillow-heif, and pytesseract to requirements.txt to enable this.")
    elif not _tesseract_ok():
        st.warning(
            "Tesseract isn't installed. On Streamlit Cloud add a packages.txt containing 'tesseract-ocr'. "
            "On Windows install the UB Mannheim build, then restart the app."
        )
    else:
        mode_label = st.radio(
            "Text location",
            ["Bottom banner (Context Camera)", "Whole photo"],
            horizontal=True,
        )
        mode = "banner" if mode_label.startswith("Bottom") else "whole"
        band = 0.08

        up = st.file_uploader(
            "Upload photos or zip files",
            type=["jpg", "jpeg", "png", "heic", "heif", "zip"],
            accept_multiple_files=True,
            key="photo_up",
        )
        if up:
            IMG_EXT = (".jpg", ".jpeg", ".png", ".heic", ".heif")
            items = []
            for f in up:
                if f.name.lower().endswith(".zip"):
                    src = f.name[:-4]
                    try:
                        with zipfile.ZipFile(io.BytesIO(f.getvalue())) as z:
                            for info in z.infolist():
                                nm = info.filename
                                base = nm.split("/")[-1]
                                if info.is_dir() or "__MACOSX" in nm or base.startswith("."):
                                    continue
                                if nm.lower().endswith(IMG_EXT):
                                    items.append((base, z.read(info), src))
                    except Exception as e:
                        st.error(f"{f.name}: {e}")
                else:
                    items.append((f.name, f.getvalue(), None))

            if not items:
                st.info("No images found in the upload.")
            else:
                with st.spinner(f"Reading {len(items)} photos..."):
                    with ThreadPoolExecutor(max_workers=4) as ex:
                        results = list(ex.map(lambda it: ocr_photo(it[1], mode, band), items))
                rows = [
                    {"File": items[i][0], "Label": results[i][0], "Date": results[i][1]}
                    for i in range(len(items))
                ]

                if st.checkbox("Auto-fix odd labels using the rest of the batch", value=True) and len(rows) > 1:
                    fixed = clean_labels([r["Label"] for r in rows])
                    for r, fl in zip(rows, fixed):
                        r["Label"] = fl

                wps = st.session_state.get("gpx_waypoints", [])
                if wps and st.checkbox("Name photos from nearest GPX waypoint (uses photo GPS)", value=False):
                    for pos, (nm, data, _s) in enumerate(items):
                        gps = _exif_gps(data)
                        if not gps:
                            continue
                        best, bestd = None, 1e12
                        for wlat, wlon, wname in wps:
                            dd = _dist_m(gps, (wlat, wlon))
                            if dd < bestd:
                                bestd, best = dd, wname
                        if best and bestd <= 100:
                            rows[pos]["Label"] = f"{_safe(best)}_{rows[pos]['Label']}"

                st.caption("Check the label and date, edit anything the OCR got wrong.")
                edited = st.data_editor(
                    pd.DataFrame(rows),
                    use_container_width=True,
                    num_rows="fixed",
                    key="photo_editor",
                    column_config={
                        "File": st.column_config.TextColumn(disabled=True),
                        "Label": st.column_config.TextColumn(),
                        "Date": st.column_config.TextColumn(help="YYYY-MM-DD. Blank goes to a 'nodate' folder."),
                    },
                )

                zip_sources = {s for _, _, s in items if s}
                multi = len(zip_sources) >= 2
                used = set()
                buf = io.BytesIO()
                with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
                    for pos, (_, r) in enumerate(edited.iterrows()):
                        src = items[pos][2]
                        date = str(r["Date"]).strip() or "nodate"
                        label = _safe(r["Label"])
                        if multi and src:
                            srcf = re.sub(r'[\\/:*?"<>|]+', "_", src).strip() or "zip"
                            folder = f"{srcf}/{date}"
                        else:
                            folder = date
                        base = f"{folder}/{label}"
                        name = base + ".jpg"
                        k = 2
                        while name in used:
                            name = f"{base}_{k}.jpg"
                            k += 1
                        used.add(name)
                        z.writestr(name, _jpeg_cached(items[pos][1]))
                st.download_button(
                    "Download zip", buf.getvalue(),
                    file_name="photos_by_date.zip", mime="application/zip",
                )
