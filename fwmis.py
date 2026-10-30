"""
FWMIS tab for the AiM Field Assistant.

Reuses the helpers and the embedded FWMIS template from fwmis_loadform_app.py
(kept in the same folder, unchanged). Only the part of that file before its UI
is executed, then render_fwmis() draws the FWMIS tab.

You type the project name and the observations; everything else fills itself:
description, datum (NAD 83), confidence (Confirmed), sensitive (N),
measurement source (GPS for lat/long rows, Derived for ATS rows),
sample numbers, count sample section for seen/heard animals, and the feature
section for signs. Also writes the submission email.
"""

import base64
import datetime as _dt
import urllib.parse
from pathlib import Path

import pandas as pd
import streamlit as st

# ---- Load helpers + embedded template from the standalone app --------------
_PATH = Path(__file__).with_name("fwmis_loadform_app.py")
_SRC = _PATH.read_text(encoding="utf-8")
_HELPERS = _SRC.split("st.set_page_config", 1)[0]   # cut off the standalone UI
_NS = {"__file__": str(_PATH), "__name__": "fwmis_loadform_app"}
exec(compile(_HELPERS, str(_PATH), "exec"), _NS)
for _k, _v in _NS.items():                          # pull in helpers + constants only
    if not _k.startswith("__"):
        globals()[_k] = _v

TEMPLATE = base64.b64decode(TEMPLATE_B64)  # noqa: F821  (from _NS)

OBS_COLS = ["Species", "Observation", "Date", "Lat, Long", "ATS Start", "ATS End", "Comments"]
SEEN, HEARD = "Seen (visual)", "Heard (call/song)"

# Count Sample columns (seen / heard animals)
COUNT_COLS = {
    "sample_no": ("BA", None),
    "start": ("BB", 69),
    "end": ("BC", 69),
    "species": ("BD", 68),
    "age": ("BG", 68),
    "gender": ("BH", 68),
    "confidence": ("BL", 68),
    "evidence": ("BM", 68),
    "comments": ("BO", 71),
}
# Feature Sample columns (signs)
FEATURE_COLS = {
    "feature": ("EC", 68),
    "date": ("EE", 69),
    "confidence": ("EF", 82),
    "species": ("EG", 68),
    "comments": ("EH", 71),
}
# Survey and location columns
BASE_COLS = {
    "survey_type": ("A", 68),
    "start": ("B", 69),
    "end": ("C", 69),
    "sensitive": ("E", 68),
    "loc_type": ("S", 73),
    "gis_source": ("U", 68),
    "datum": ("V", 68),
    "in_alberta": ("X", 68),
    "lat": ("Y", 84),
    "lon": ("Z", 84),
}
ATS_COLS = {
    "s_": {"lsd": "AD", "quarter": "AE", "section": "AF", "township": "AG", "range": "AH", "meridian": "AI"},
    "e_": {"lsd": "AO", "quarter": "AP", "section": "AQ", "township": "AR", "range": "AS", "meridian": "AT"},
}

NAME_HELP = ("Format: client, project, then \"Wildlife Sweep\". "
             "Example: Keyera GWD 31640 and 31670 Wildlife Sweep. "
             "The description and email fill themselves from this and the dates.")
LICENCE_HELP = ("The research licence number is for species wildlife surveys, typically where "
                "playback surveys or capture occurs. Leave as NA for a standard sweep.")
OBS_HELP = ("Seen = you saw the animal. Heard = call or song. "
            "Anything else is a sign (tracks, scat, nest, etc). Blank counts as Seen.")
DATE_HELP = "Leave blank to use the sweep date. Only fill if this one was on a different day."
ATS_HELP = ("Fill only if this row is in a different spot than the default ATS above. "
            "Format NE-16-75-10-W6M.")


# ---------------------------------------------------------------------------
# Text helpers
# ---------------------------------------------------------------------------
def _fmt_date(d):
    return f"{d:%B} {d.day}, {d.year}"


def _date_phrase(dates):
    """'October 26, 2026' or 'October 26 to 28, 2026' or 'October 30 to November 2, 2026'."""
    lo, hi = min(dates), max(dates)
    if lo == hi:
        return _fmt_date(lo)
    if lo.year == hi.year and lo.month == hi.month:
        return f"{lo:%B} {lo.day} to {hi.day}, {hi.year}"
    if lo.year == hi.year:
        return f"{lo:%B} {lo.day} to {hi:%B} {hi.day}, {hi.year}"
    return f"{_fmt_date(lo)} to {_fmt_date(hi)}"


def _project_phrase(name):
    """'Keyera GWD 31640 and 31670 Wildlife Sweep' -> "Keyera's GWD 31640 and 31670"."""
    s = re.sub(r"\s*wildlife\s+sweep\s*$", "", name.strip(), flags=re.I).strip()  # noqa: F821
    if not s:
        return name.strip()
    parts = s.split(" ", 1)
    client = parts[0]
    owner = client + ("'" if client.lower().endswith("s") else "'s")
    return f"{owner} {parts[1]}" if len(parts) > 1 else client


def _extra_lists():
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(TEMPLATE), read_only=True, data_only=True)  # noqa: F821
    rows = list(wb[CODES_SHEET].iter_rows(values_only=True))  # noqa: F821
    hdr = rows[0]
    out = {}
    for name in ["inv_EVIDENCE", "inv_AGE_GROUP"]:
        if name in hdr:
            i = hdr.index(name)
            out[name] = [r[i] for r in rows[1:] if i < len(r) and r[i]]
        else:
            out[name] = []
    wb.close()
    return out


# ---------------------------------------------------------------------------
# Loadform writer
# ---------------------------------------------------------------------------
def _build_loadform(project, survey, rows):
    zin = zipfile.ZipFile(io.BytesIO(TEMPLATE))  # noqa: F821
    paths = sheet_paths(zin)  # noqa: F821
    sst = SharedStrings(zin.read("xl/sharedStrings.xml"))  # noqa: F821

    proj = SheetXml(zin.read(paths[PROJECT_SHEET]))  # noqa: F821
    proj.clear_values_from(FIRST_DATA_ROW)  # noqa: F821
    for key, (col, style) in PROJECT_CELLS.items():  # noqa: F821
        proj.set(f"{col}{FIRST_DATA_ROW}", project.get(key), sst, style)  # noqa: F821

    surv = SheetXml(zin.read(paths[SURVEY_SHEET]))  # noqa: F821
    surv.clear_values_from(FIRST_DATA_ROW)  # noqa: F821

    def put(col, rn, value, style):
        surv.set(f"{col}{rn}", value, sst, style)

    for i, r in enumerate(rows):
        rn = FIRST_DATA_ROW + i  # noqa: F821
        base = {
            "survey_type": survey["survey_type"],
            "start": excel_serial(survey["start"]),  # noqa: F821
            "end": excel_serial(survey["end"]),  # noqa: F821
            "sensitive": "N",
            "loc_type": r["loc_type"],
            "gis_source": r["gis_source"],
            "datum": survey["datum"],
            "in_alberta": "Y",
            "lat": round(float(r["lat"]), 6) if r.get("lat") is not None else None,
            "lon": round(float(r["lon"]), 6) if r.get("lon") is not None else None,
        }
        for key, (col, style) in BASE_COLS.items():
            put(col, rn, base[key], style)

        for prefix, ats in (("s_", r.get("ats_start")), ("e_", r.get("ats_end"))):
            for field, col in ATS_COLS[prefix].items():
                put(col, rn, (ats or {}).get(field), 74)

        put(COUNT_COLS["sample_no"][0], rn, i + 1, None)   # sample # on every row

        day = excel_serial(r["date"])  # noqa: F821
        if r["kind"] == "count":
            vals = {"start": day, "end": day, "species": r["species"], "age": survey["age"],
                    "gender": "U", "confidence": survey["confidence"], "evidence": r["evidence"],
                    "comments": r["comment"] or None}
            for key, (col, style) in COUNT_COLS.items():
                if key != "sample_no":
                    put(col, rn, vals[key], style)
        else:
            vals = {"feature": r["feature"], "date": day, "confidence": survey["confidence"],
                    "species": r["species"], "comments": r["comment"] or None}
            for key, (col, style) in FEATURE_COLS.items():
                put(col, rn, vals[key], style)

    changed = {
        paths[PROJECT_SHEET]: proj.to_bytes(),
        paths[SURVEY_SHEET]: surv.to_bytes(),
        "xl/sharedStrings.xml": sst.to_bytes(),
    }
    out = io.BytesIO()  # noqa: F821
    with zipfile.ZipFile(out, "w") as zout:  # noqa: F821
        for info in zin.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=info.date_time)  # noqa: F821
            zi.compress_type = info.compress_type
            zi.external_attr = info.external_attr
            zout.writestr(zi, changed.get(info.filename, zin.read(info.filename)))
    return out.getvalue()


# ---------------------------------------------------------------------------
# ATS picker
# ---------------------------------------------------------------------------
def _ats_picker(prefix, label):
    st.markdown(label)
    c = st.columns(5)
    q = c[0].selectbox("Quarter / LSD", [""] + ["NE", "NW", "SE", "SW"] + [str(n) for n in range(1, 17)],
                       key=f"{prefix}_q")
    sec = c[1].selectbox("Section", [""] + [str(n) for n in range(1, 37)], key=f"{prefix}_sec")
    twp = c[2].selectbox("Township", [""] + [str(n) for n in range(1, 127)], key=f"{prefix}_twp")
    rge = c[3].selectbox("Range", [""] + [str(n) for n in range(1, 31)], key=f"{prefix}_rge")
    mer = c[4].selectbox("Meridian", ["", "W4M", "W5M", "W6M"], key=f"{prefix}_mer")
    if q and sec and twp and rge and mer:
        return f"{q}-{sec}-{twp}-{rge}-{mer}"
    return ""


# ---------------------------------------------------------------------------
# UI
# ---------------------------------------------------------------------------
def render_fwmis():
    st.subheader("FWMIS Loadform Builder")
    st.caption("Type the project name and the observations, then download the loadform and copy the email.")

    lists = read_code_lists(TEMPLATE)  # noqa: F821
    extra = _extra_lists()
    SPECIES = lists["tax_TAXONOMIC"]
    SIGNS = [f for f in lists["inv_SITFEATURE"] if code_of(f) != "AUDI"]  # noqa: F821
    OBS_OPTIONS = [SEEN, HEARD] + SIGNS
    SITE = pick_option(lists["inv_LOCTYPE"], "SITE")  # noqa: F821
    TRANSECT = pick_option(lists["inv_LOCTYPE"], "TRANSECT")  # noqa: F821
    GPS = pick_option(lists["inv_GISSOURCE"], "GPS")  # noqa: F821
    DERIVED = pick_option(lists["inv_GISSOURCE"], "DERIVED")  # noqa: F821
    DATUM = pick_option(lists["inv_GISDATUM"], "NAD 83")  # noqa: F821
    CONFIDENCE = pick_option(lists["inv_CONFDNTLVL"], "CONFIRMED")  # noqa: F821
    OBSERVED = pick_option(extra["inv_EVIDENCE"], "OBSERVED") or "Observed -OBSERVED"  # noqa: F821
    HEARD_EV = pick_option(extra["inv_EVIDENCE"], "HEARD") or "Heard -HEARD"  # noqa: F821
    AGE_U = pick_option(extra["inv_AGE_GROUP"], "U") or "Unknown -U"  # noqa: F821
    AREASEARCH = pick_option(lists["inv_SURVEYTYPE"], "AREASEARCH") or lists["inv_SURVEYTYPE"][0]  # noqa: F821

    # --- Project -----------------------------------------------------------
    st.markdown("**Project**")
    c1, c2 = st.columns([2, 1])
    proj_name = c1.text_input("Project name", key="fwmis_name", help=NAME_HELP,
                              placeholder="Client Name Project Name Wildlife Sweep")
    licence = c2.text_input("Research licence number", value="NA", key="fwmis_lic", help=LICENCE_HELP)

    c1, c2 = st.columns(2)
    sweep_date = c1.date_input("Sweep date", format="YYYY-MM-DD", key="fwmis_date",
                               help="Used for every row unless you set a different date on a row.")
    survey_type = c2.selectbox("Survey type", lists["inv_SURVEYTYPE"],
                               index=lists["inv_SURVEYTYPE"].index(AREASEARCH), key="fwmis_stype")

    # --- Default location ---------------------------------------------------
    st.markdown("**Default ATS** (used for every row that has no Lat, Long or ATS of its own)")
    common_s = _ats_picker("fwmis_s", "Start")
    two_q = st.checkbox("Sweep crosses into a second quarter (add an ATS End)", key="fwmis_two_q")
    common_e = _ats_picker("fwmis_e", "End") if two_q else ""
    if common_s:
        st.caption(f"Default ATS: {common_s}" + (f" to {common_e}" if common_e else ""))

    # --- Observations ------------------------------------------------------
    st.markdown("**Observations** (one row each)")
    st.caption("Only Species is required. Observation blank = Seen. Date blank = sweep date. "
               "Location blank = default ATS.")
    if "fwmis_rows" not in st.session_state:
        st.session_state["fwmis_rows"] = pd.DataFrame(
            {"Species": [None] * 8, "Observation": [None] * 8, "Date": [None] * 8,
             "Lat, Long": [""] * 8, "ATS Start": [""] * 8, "ATS End": [""] * 8, "Comments": [""] * 8})

    edited = st.data_editor(
        st.session_state["fwmis_rows"], num_rows="dynamic", hide_index=True, key="fwmis_editor",
        column_config={
            "Species": st.column_config.SelectboxColumn("Species", options=SPECIES, width="large"),
            "Observation": st.column_config.SelectboxColumn("Observation", options=OBS_OPTIONS,
                                                            width="medium", help=OBS_HELP),
            "Date": st.column_config.DateColumn("Date", format="YYYY-MM-DD", help=DATE_HELP),
            "Lat, Long": st.column_config.TextColumn("Lat, Long", help="e.g. 55.1234, -119.1234"),
            "ATS Start": st.column_config.TextColumn("ATS Start", help=ATS_HELP),
            "ATS End": st.column_config.TextColumn("ATS End", help=ATS_HELP),
            "Comments": st.column_config.TextColumn("Comments", width="large"),
        })

    # --- Build rows --------------------------------------------------------
    rows, errors = [], []
    for i, r in edited.reset_index(drop=True).iterrows():
        sp = r.get("Species")
        obs = r.get("Observation")
        ll_text = str(r.get("Lat, Long") or "").strip()
        ats_s_text = str(r.get("ATS Start") or "").strip()
        ats_e_text = str(r.get("ATS End") or "").strip()
        comment = str(r.get("Comments") or "").strip()
        if not sp and not obs and not ll_text and not ats_s_text and not ats_e_text and not comment:
            continue  # blank row
        n = i + 1
        if not sp:
            errors.append(f"Row {n}: pick a species.")
            continue

        day = to_date(r.get("Date")) or sweep_date  # noqa: F821
        obs = obs or SEEN
        row = {"row": n, "date": day, "species": sp, "comment": comment, "lat": None, "lon": None}
        if obs in (SEEN, HEARD):
            row.update(kind="count", evidence=HEARD_EV if obs == HEARD else OBSERVED)
        else:
            row.update(kind="feature", feature=obs)
            if code_of(obs) == "OTHR" and not comment:  # noqa: F821
                errors.append(f"Row {n}: 'Other' needs a comment saying what it was.")

        ll = parse_latlon(ll_text) if ll_text else None  # noqa: F821
        if ll_text and not ll:
            errors.append(f"Row {n}: Lat, Long doesn't look right. Use e.g. 55.1234, -119.1234.")
            continue
        if ll:
            lat, lon = ll
            row.update(lat=lat, lon=lon, loc_type=SITE, gis_source=GPS)
            if not (ALBERTA_LAT[0] <= lat <= ALBERTA_LAT[1]) or not (ALBERTA_LON[0] <= lon <= ALBERTA_LON[1]):  # noqa: F821
                errors.append(f"Row {n}: {lat}, {lon} is outside Alberta (longitude should be negative).")
        else:
            if not ats_s_text and not ats_e_text:
                ats_s_text, ats_e_text = common_s, common_e
            s_ats = parse_ats(ats_s_text)  # noqa: F821
            e_ats = parse_ats(ats_e_text)  # noqa: F821
            if s_ats is None and e_ats is None:
                if ats_s_text or ats_e_text:
                    errors.append(f"Row {n}: ATS doesn't look right. Use e.g. NE-16-75-10-W6M.")
                else:
                    errors.append(f"Row {n}: needs a location. Fill the default ATS above, "
                                  "or a Lat, Long / ATS on this row.")
                continue
            if s_ats is None:
                s_ats, e_ats = e_ats, None
            if e_ats == s_ats:
                e_ats = None
            row.update(ats_start=ats_to_fwmis(s_ats, lists), ats_end=ats_to_fwmis(e_ats, lists),  # noqa: F821
                       loc_type=TRANSECT if e_ats else SITE, gis_source=DERIVED)
        rows.append(row)

    if not proj_name.strip():
        errors.append("Add a project name.")
    if not rows:
        errors.append("Add at least one observation.")

    # --- Auto text ---------------------------------------------------------
    dates = [r["date"] for r in rows] or [sweep_date]
    date_txt = _date_phrase(dates)
    phrase = _project_phrase(proj_name) if proj_name.strip() else "the project"
    verb = "completed on" if min(dates) == max(dates) else "completed from"
    proj_desc = f"A preconstruction wildlife sweep was {verb} {date_txt}, for {phrase}."

    if proj_name.strip():
        st.caption(f"Description: {proj_desc}")

    for e in errors:
        st.error(e)

    file_bytes = None
    if not errors:
        file_bytes = _build_loadform(
            project={"name": proj_name.strip(), "description": proj_desc, "licence": licence.strip()},
            survey={"survey_type": survey_type, "start": min(dates), "end": max(dates),
                    "datum": DATUM, "confidence": CONFIDENCE, "age": AGE_U},
            rows=rows,
        )

    safe = re.sub(r"[^A-Za-z0-9]+", "_", proj_name).strip("_") or "Project"  # noqa: F821
    st.download_button(
        f"Download loadform ({len(rows)} rows)" if file_bytes else "Download loadform",
        data=file_bytes or b"",
        file_name=f"FWMIS_LOADFORM_{safe}_{min(dates):%Y%m%d}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        disabled=file_bytes is None,
        type="primary",
        key="fwmis_dl",
    )

    # --- Email -------------------------------------------------------------
    if proj_name.strip():
        st.markdown("**Submission email**")
        subject = f"FWMIS Loadform Submission - {proj_name.strip()}"
        body = (
            "Hello,\n\n"
            "Please find attached the completed FWMIS load form for wildlife observations recorded "
            f"during the {date_txt}, wildlife sweep for {phrase}.\n\n"
            "This submission is being provided to you as the appropriate contact identified through the "
            "Alberta Environment and Protected Areas Fish and Wildlife Contact Map using the Wildlife "
            "Contact District layer.\n\n"
            "Please let me know if any additional information is required.\n\n"
            "Kind regards,"
        )
        st.caption("Subject")
        st.code(subject, language=None)
        st.caption("Body")
        st.code(body, language=None)
        mailto = "mailto:?subject=" + urllib.parse.quote(subject) + "&body=" + urllib.parse.quote(body)
        st.link_button("Open in email app", mailto)
        st.caption("Remember to attach the downloaded loadform.")
